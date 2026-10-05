"""
Tests for Phase 4: PDF/DOCX/TXT/MD extraction and structured chunks.
"""

from __future__ import annotations

from io import BytesIO

import fitz
import pytest
from docx import Document
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import ApplicationVersion, DocumentChunk, GuidelineVersion


async def create_project(client: AsyncClient) -> str:
    response = await client.post("/api/projects", json={"name": "Extraction Project"})
    assert response.status_code == 201
    return response.json()["id"]


def make_pdf(text_by_page: list[str]) -> bytes:
    document = fitz.open()
    for text in text_by_page:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def make_docx() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("Eligibility", level=1)
    document.add_paragraph("Applicants must be nonprofit organizations.")
    document.add_paragraph("A budget narrative is recommended.")
    document.save(buffer)
    return buffer.getvalue()


class TestTextExtraction:
    """Successful extraction for every supported format."""

    @pytest.mark.asyncio
    async def test_txt_upload_extracts_raw_text_and_chunks(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
    ):
        project_id = await create_project(client)
        response = await client.post(
            f"/api/projects/{project_id}/guidelines",
            files={
                "file": (
                    "guideline.txt",
                    b"Eligibility\n\nApplicants must include audited financials.",
                    "text/plain",
                )
            },
        )

        assert response.status_code == 201
        version_id = response.json()["version"]["id"]
        assert response.json()["version"]["total_chunks"] == 2

        version = await db_session.get(GuidelineVersion, version_id)
        assert version is not None
        assert "audited financials" in version.raw_text
        assert version.total_chunks == 2

        chunks = (
            await db_session.execute(
                select(DocumentChunk)
                .where(DocumentChunk.guideline_version_id == version_id)
                .order_by(DocumentChunk.chunk_index)
            )
        ).scalars().all()
        assert [chunk.chunk_index for chunk in chunks] == [0, 1]
        assert chunks[1].text == "Applicants must include audited financials."
        assert chunks[1].char_start is not None
        assert chunks[1].token_count == 5

    @pytest.mark.asyncio
    async def test_markdown_upload_tracks_heading_metadata(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
    ):
        project_id = await create_project(client)
        response = await client.post(
            f"/api/projects/{project_id}/applications",
            files={
                "file": (
                    "application.md",
                    b"# Timeline\n\nThe project runs from January to December 2027.",
                    "text/markdown",
                )
            },
        )

        assert response.status_code == 201
        version_id = response.json()["version"]["id"]
        chunks = (
            await db_session.execute(
                select(DocumentChunk)
                .where(DocumentChunk.application_version_id == version_id)
                .order_by(DocumentChunk.chunk_index)
            )
        ).scalars().all()

        assert len(chunks) == 2
        assert chunks[0].section_heading == "Timeline"
        assert chunks[1].section_heading == "Timeline"

    @pytest.mark.asyncio
    async def test_pdf_upload_extracts_page_numbers(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
    ):
        project_id = await create_project(client)
        pdf_bytes = make_pdf(
            [
                "Eligibility\nApplicants must be registered nonprofits.",
                "Budget\nApplicants must provide a detailed budget.",
            ]
        )

        response = await client.post(
            f"/api/projects/{project_id}/guidelines",
            files={"file": ("guideline.pdf", pdf_bytes, "application/pdf")},
        )

        assert response.status_code == 201
        data = response.json()
        assert data["version"]["total_pages"] == 2
        assert data["version"]["total_chunks"] == 4

        chunks = (
            await db_session.execute(
                select(DocumentChunk)
                .where(DocumentChunk.guideline_version_id == data["version"]["id"])
                .order_by(DocumentChunk.chunk_index)
            )
        ).scalars().all()
        assert [chunk.page_number for chunk in chunks] == [1, 1, 2, 2]
        assert "registered nonprofits" in chunks[1].text

    @pytest.mark.asyncio
    async def test_docx_upload_extracts_paragraphs_and_headings(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
    ):
        project_id = await create_project(client)
        response = await client.post(
            f"/api/projects/{project_id}/applications",
            files={
                "file": (
                    "draft.docx",
                    make_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        assert response.status_code == 201
        version_id = response.json()["version"]["id"]
        version = await db_session.get(ApplicationVersion, version_id)
        assert version is not None
        assert version.total_chunks == 3
        assert "nonprofit organizations" in version.raw_text

        chunks = (
            await db_session.execute(
                select(DocumentChunk)
                .where(DocumentChunk.application_version_id == version_id)
                .order_by(DocumentChunk.chunk_index)
            )
        ).scalars().all()
        assert chunks[0].section_heading == "Eligibility"
        assert chunks[1].section_heading == "Eligibility"


class TestExtractionFailures:
    """Extraction failures should be clear and store no partial version."""

    @pytest.mark.asyncio
    async def test_corrupt_pdf_returns_extraction_error(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
    ):
        project_id = await create_project(client)
        response = await client.post(
            f"/api/projects/{project_id}/guidelines",
            files={"file": ("broken.pdf", b"not a real pdf", "application/pdf")},
        )

        assert response.status_code == 422
        assert response.json()["error"]["code"] == "PDF_EXTRACTION_FAILED"

        versions = (await db_session.execute(select(GuidelineVersion))).scalars().all()
        assert versions == []

    @pytest.mark.asyncio
    async def test_image_only_pdf_returns_no_ocr_error(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
    ):
        project_id = await create_project(client)
        image_only_pdf = make_pdf([""])

        response = await client.post(
            f"/api/projects/{project_id}/guidelines",
            files={"file": ("image-only.pdf", image_only_pdf, "application/pdf")},
        )

        assert response.status_code == 422
        assert response.json()["error"]["code"] == "NO_EXTRACTABLE_TEXT"
        assert "OCR is not supported" in response.json()["error"]["message"]

        versions = (await db_session.execute(select(GuidelineVersion))).scalars().all()
        assert versions == []

    @pytest.mark.asyncio
    async def test_corrupt_docx_returns_extraction_error(
        self,
        client: AsyncClient,
    ):
        project_id = await create_project(client)
        response = await client.post(
            f"/api/projects/{project_id}/applications",
            files={
                "file": (
                    "broken.docx",
                    b"not a real docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        assert response.status_code == 422
        assert response.json()["error"]["code"] == "DOCX_EXTRACTION_FAILED"

    @pytest.mark.asyncio
    async def test_invalid_text_encoding_returns_decode_error(
        self,
        client: AsyncClient,
    ):
        project_id = await create_project(client)
        response = await client.post(
            f"/api/projects/{project_id}/applications",
            files={"file": ("draft.txt", b"\xff\xfe\x00\x00", "text/plain")},
        )

        assert response.status_code == 422
        assert response.json()["error"]["code"] == "TEXT_DECODE_FAILED"
